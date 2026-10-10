"""Generate a small, reproducible sales workbook from a CSV export."""

from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


REQUIRED = {"order_id", "date", "product", "quantity", "unit_price"}
MONEY_FORMAT = '#,##0.00 "₽"'
HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass(frozen=True)
class Order:
    order_id: str
    day: date
    product: str
    quantity: int
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return self.quantity * self.unit_price


def safe_text(value: str) -> str:
    """Keep imported text from becoming an Excel formula."""
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def read_orders(path: Path) -> tuple[list[Order], list[dict]]:
    orders: list[Order] = []
    errors: list[dict] = []
    seen_ids: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not REQUIRED.issubset(reader.fieldnames):
            missing = sorted(REQUIRED - set(reader.fieldnames or []))
            raise ValueError(f"не хватает колонок: {', '.join(missing)}")
        for row_number, row in enumerate(reader, start=2):
            try:
                if None in row:
                    raise ValueError("колонок больше, чем в заголовке")
                order_id = (row["order_id"] or "").strip()
                product = (row["product"] or "").strip()
                if not order_id or not product:
                    raise ValueError("пустой номер заказа или товар")
                if order_id in seen_ids:
                    raise ValueError("повторный номер заказа")
                day = date.fromisoformat((row["date"] or "").strip())
                quantity = int((row["quantity"] or "").strip())
                price = Decimal((row["unit_price"] or "").strip())
                if quantity <= 0 or not price.is_finite() or price < 0:
                    raise ValueError("количество должно быть положительным, цена — неотрицательной")
                if price.as_tuple().exponent < -2:
                    raise ValueError("у цены должно быть не более двух знаков после запятой")
                orders.append(Order(order_id, day, product, quantity, price))
                seen_ids.add(order_id)
            except (ValueError, InvalidOperation, TypeError) as exc:
                errors.append({"row": row_number, "order_id": row.get("order_id", ""), "reason": str(exc)})
    return orders, errors


def summarize(orders: list[Order], errors: list[dict]) -> dict:
    total = sum((order.amount for order in orders), Decimal("0"))
    average = total / len(orders) if orders else Decimal("0")
    return {
        "valid_orders": len(orders),
        "invalid_rows": len(errors),
        "revenue_rub": f"{total:.2f}",
        "average_order_rub": f"{average:.2f}",
    }


def style_table(sheet, widths: dict[str, int], money_columns: tuple[int, ...] = ()) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for column in money_columns:
            row[column - 1].number_format = MONEY_FORMAT


def make_workbook(orders: list[Order], errors: list[dict]) -> Workbook:
    summary = summarize(orders, errors)
    book = Workbook()
    overview = book.active
    overview.title = "Итоги"
    overview.append(["Показатель", "Значение"])
    overview.append(["Выручка, ₽", float(Decimal(summary["revenue_rub"]))])
    overview.append(["Заказов", summary["valid_orders"]])
    overview.append(["Средний чек, ₽", float(Decimal(summary["average_order_rub"]))])
    overview.append(["Строк с ошибками", summary["invalid_rows"]])
    style_table(overview, {"A": 27, "B": 20})
    overview["B2"].number_format = MONEY_FORMAT
    overview["B4"].number_format = MONEY_FORMAT

    by_day: dict[date, dict] = defaultdict(lambda: {"count": 0, "amount": Decimal("0")})
    by_product: dict[str, dict] = defaultdict(lambda: {"quantity": 0, "amount": Decimal("0")})
    for order in orders:
        by_day[order.day]["count"] += 1
        by_day[order.day]["amount"] += order.amount
        by_product[order.product]["quantity"] += order.quantity
        by_product[order.product]["amount"] += order.amount

    days = book.create_sheet("По дням")
    days.append(["Дата", "Заказов", "Выручка, ₽"])
    for day, values in sorted(by_day.items()):
        days.append([day, values["count"], float(values["amount"])])
        days.cell(days.max_row, 1).number_format = "DD.MM.YYYY"
    style_table(days, {"A": 17, "B": 15, "C": 20}, money_columns=(3,))
    if by_day:
        chart = BarChart()
        chart.title = "Выручка по дням"
        chart.y_axis.title = "Рубли"
        chart.x_axis.title = "Дата"
        chart.add_data(Reference(days, min_col=3, min_row=1, max_row=days.max_row), titles_from_data=True)
        chart.set_categories(Reference(days, min_col=1, min_row=2, max_row=days.max_row))
        chart.width = 18
        chart.height = 9
        days.add_chart(chart, "E2")

    products = book.create_sheet("По товарам")
    products.append(["Товар", "Продано штук", "Выручка, ₽"])
    for product, values in sorted(by_product.items(), key=lambda pair: (-pair[1]["amount"], pair[0])):
        products.append([safe_text(product), values["quantity"], float(values["amount"])])
    style_table(products, {"A": 30, "B": 19, "C": 20}, money_columns=(3,))

    detail = book.create_sheet("Заказы")
    detail.append(["Номер заказа", "Дата", "Товар", "Количество", "Цена, ₽", "Сумма, ₽"])
    for order in orders:
        detail.append([safe_text(order.order_id), order.day, safe_text(order.product),
                       order.quantity, float(order.unit_price), float(order.amount)])
        detail.cell(detail.max_row, 2).number_format = "DD.MM.YYYY"
    style_table(detail, {"A": 20, "B": 17, "C": 30, "D": 16, "E": 16, "F": 18},
                money_columns=(5, 6))

    problems = book.create_sheet("Ошибки")
    problems.append(["Строка CSV", "Номер заказа", "Причина"])
    for error in errors:
        problems.append([error["row"], safe_text(error["order_id"] or ""), error["reason"]])
    style_table(problems, {"A": 16, "B": 20, "C": 65})
    return book


def main() -> int:
    parser = argparse.ArgumentParser(description="Построить Excel-отчёт по продажам")
    parser.add_argument("input", type=Path, help="CSV с заказами")
    parser.add_argument("--out", type=Path, default=Path("output"), help="папка результата")
    args = parser.parse_args()
    try:
        orders, errors = read_orders(args.input)
        args.out.mkdir(parents=True, exist_ok=True)
        make_workbook(orders, errors).save(args.out / "report.xlsx")
        summary = summarize(orders, errors)
        (args.out / "summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    except (OSError, UnicodeError, ValueError, csv.Error) as exc:
        parser.exit(1, f"Ошибка: {exc}\n")
    print(f"Готово: {summary['valid_orders']} заказов, выручка {summary['revenue_rub']} ₽. Папка: {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
