import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from report import make_workbook, read_orders, safe_text, summarize  # noqa: E402


class ReportTests(unittest.TestCase):
    def test_sample_totals_and_invalid_row(self):
        source = Path(__file__).resolve().parents[1] / "examples/orders.csv"
        orders, errors = read_orders(source)
        self.assertEqual(summarize(orders, errors), {
            "valid_orders": 6, "invalid_rows": 1,
            "revenue_rub": "3322.50", "average_order_rub": "553.75",
        })
        self.assertEqual(orders[0].amount, Decimal("900.00"))
        self.assertEqual(errors[0]["row"], 7)

    def test_workbook_contains_breakdowns_and_chart(self):
        source = Path(__file__).resolve().parents[1] / "examples/orders.csv"
        orders, errors = read_orders(source)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "report.xlsx"
            make_workbook(orders, errors).save(path)
            book = load_workbook(path)
            self.assertEqual(book.sheetnames, ["Итоги", "По дням", "По товарам", "Заказы", "Ошибки"])
            self.assertEqual(book["Итоги"]["B2"].value, 3322.5)
            self.assertEqual(book["По дням"]["C2"].value, 1498.5)
            self.assertEqual(len(book["По дням"]._charts), 1)
            self.assertEqual(book["Ошибки"].max_row, 2)

    def test_duplicate_order_and_formula_like_text(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "orders.csv"
            path.write_text("order_id,date,product,quantity,unit_price\n"
                            "X1,2026-09-01,=1+1,1,5.00\n"
                            "X1,2026-09-02,Кружка,1,5.00\n", encoding="utf-8")
            orders, errors = read_orders(path)
            self.assertEqual(len(orders), 1)
            self.assertEqual(errors[0]["reason"], "повторный номер заказа")
            book = make_workbook(orders, errors)
            self.assertEqual(book["Заказы"]["C2"].data_type, "s")
            self.assertEqual(book["Заказы"]["C2"].value, safe_text("=1+1"))

    def test_missing_column(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "orders.csv"
            path.write_text("order_id,date,product,quantity\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "unit_price"):
                read_orders(path)


if __name__ == "__main__":
    unittest.main()
